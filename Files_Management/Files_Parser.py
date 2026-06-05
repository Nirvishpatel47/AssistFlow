# Files_Management/Files_Parser.py
from Security.Advance_Logger import logger
from pathlib import Path
from email import policy
from email.parser import BytesParser
from bs4 import BeautifulSoup
import pymupdf4llm
import numpy as np
import openpyxl
import docx
import re
import asyncio

class FileParser:
    @staticmethod
    def sanitize_text(user_input: str) -> str:
        """
        Removes potentially dangerous characters and trims whitespace.
        Runs synchronously as it is purely a lightweight CPU operation.
        """
        if not isinstance(user_input, str):
            raise ValueError("Input must be a string")
        
        safe_text = re.sub(r"[^a-zA-Z0-9\s.,!?@#_-]", "", user_input)
        return safe_text.strip()

    @staticmethod
    def _read_file_sync(path: str) -> str:
        """Isolated helper for synchronous file read."""
        with open(path, "rb") as f:
            return f.read().decode("utf-8", errors="ignore")

    @staticmethod
    async def parse_pdf(path: str) -> str:
        try:
            loop = asyncio.get_running_loop()
            return await loop.run_in_executor(None, pymupdf4llm.to_markdown, path)
        except Exception as e:
            logger.error("FileParser.parse_pdf", e)
            return ""
    
    @staticmethod
    async def parse_docx(path: str) -> str:
        try:
            loop = asyncio.get_running_loop()
            def _read_docx():
                doc = docx.Document(path)
                return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            return await loop.run_in_executor(None, _read_docx)
        except Exception as e:
            logger.error("FileParser.parse_docx", e)
            return "" 

    @staticmethod
    async def extract_codebase(path: str) -> str:
        try:
            loop = asyncio.get_running_loop()
            return await loop.run_in_executor(None, FileParser._read_file_sync, path)
        except Exception as e:
            logger.error("FileParser.extract_codebase", e)
            return ""
    
    @staticmethod
    async def extract_email(path: str) -> str:
        try:
            loop = asyncio.get_running_loop()
            def _parse_email():
                with open(path, "rb") as f:
                    return BytesParser(policy=policy.default).parse(f)
            
            msg = await loop.run_in_executor(None, _parse_email)
            body = None
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    try:
                        body = part.get_content()
                        break
                    except Exception:
                        pass

            # Fallback to HTML
            if body is None:
                for part in msg.walk():
                    if part.get_content_type() == "text/html":
                        try:
                            html = part.get_content()
                            soup = BeautifulSoup(html, "lxml")
                            for tag in soup(["script", "style", "img", "meta", "head"]):
                                tag.decompose()
                            body = soup.get_text("\n")
                            break
                        except Exception:
                            pass
            if not body:
                body = ""
            
            lines = []
            append = lines.append
            for line in body.splitlines():
                line = line.strip()
                if not line:
                    continue
                lower = line.lower()
                if ("unsubscribe" in lower or "tracking" in lower or "emailopen" in lower):
                    continue
                append(line)
            return "\n".join(lines)
        except Exception as e:
            logger.error("FileParser.extract_email", e)
            return ""

    @staticmethod
    async def parse_excel(path: str) -> str:
        try:
            loop = asyncio.get_running_loop()
            if path.lower().endswith(".csv"):
                def _read_csv():
                    import csv
                    local_rows = []
                    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
                        reader = csv.reader(f)
                        for row in reader:
                            values = [str(cell).strip() for cell in row if cell not in (None, "")]
                            if values:
                                local_rows.append(" | ".join(values))
                    return "\n".join(local_rows).strip()
                
                return await loop.run_in_executor(None, _read_csv)

            def _read_excel():
                local_rows = []
                wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
                for sheet in wb.worksheets:
                    local_rows.append(f"\nSheet: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(cell).strip() for cell in row if cell is not None]
                        if values:
                            local_rows.append(" | ".join(values))
                return "\n".join(local_rows).strip()

            return await loop.run_in_executor(None, _read_excel)
        except Exception as e:
            logger.error("FileParser.parse_excel", e)
            return ""
        
class ParseFile:
    PDF = ".pdf"
    DOCX = ".docx"
    EML = ".eml"
    XL = frozenset({".xlsx", ".xls", ".csv"})
    CODE_EXTENSIONS = frozenset({
        ".py", ".pyi", ".js", ".jsx", ".mjs", ".cjs", ".ts", ".tsx", ".java",
        ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".hxx", ".cs", ".go",
        ".rs", ".php", ".rb", ".swift", ".kt", ".kts", ".scala", ".dart", ".r",
        ".lua", ".pl", ".sh", ".bash", ".zsh", ".ps1", ".sql", ".html", ".htm",
        ".css", ".scss", ".sass", ".less", ".json", ".xml", ".yaml", ".yml",
        ".toml", ".ini", ".cfg", ".env", ".md", ".txt", ".dockerfile", ".tf",
        ".vue", ".svelte", ".ipynb"
    })

    @staticmethod
    async def parse_any_file(path: str) -> str:
        try:
            suffix = Path(path).suffix.lower()
            if suffix == ParseFile.PDF:
                return await FileParser.parse_pdf(path)
            if suffix == ParseFile.DOCX:
                return await FileParser.parse_docx(path)
            if suffix == ParseFile.EML:
                return await FileParser.extract_email(path)
            if suffix in ParseFile.CODE_EXTENSIONS:
                return await FileParser.extract_codebase(path)
            if suffix in ParseFile.XL:
                return await FileParser.parse_excel(path)
            return await FileParser.extract_codebase(path)
        except Exception as e:
            logger.error("ParseFile.parse_any_file", e)
            return ""
        
class Chunker:
    @staticmethod
    def split_into_sentences(text: str) -> list[str]:
        """Splits raw text into sentences synchronously using regex."""
        if not text:
            return []
        sentences = re.split(r'(?<=[.!?])\s+', text)
        return [s.strip() for s in sentences if s.strip()]
    
    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1200, chunk_overlap: int = 200) -> list[str]:
        """Splits text files into overlapping sliding-window segments."""
        if not text or not text.strip():
            return []
        
        chunks = []
        start = 0
        text_len = len(text)
        
        while start < text_len:
            end = start + chunk_size
            chunks.append(text[start:end].strip())
            start += (chunk_size - chunk_overlap)
            
        return chunks

    @staticmethod
    async def chunk_text_semantically( text: str, similarity_threshold: float = 0.07, min_chunk_size: int = 800 ) -> list[str]:
        """
        Lightweight semantic-ish chunking without embeddings.

        Features:
        - No AI models
        - No numpy
        - No sklearn
        - Very fast
        - Multi-user safe
        - Async-compatible
        """

        if not text or not text.strip():
            return []

        sentences = Chunker.split_into_sentences(text)

        if len(sentences) <= 1:
            return sentences

        stopwords = {
            "the", "a", "an", "and", "or", "but",
            "is", "are", "was", "were", "to",
            "of", "in", "on", "for", "with",
            "that", "this", "it", "as", "at",
            "by", "from"
        }

        def tokenize(sentence: str) -> set[str]:
            return {
                word.strip(".,!?():;\"'").lower()
                for word in sentence.split()
                if len(word) > 2
                and word.lower() not in stopwords
            }

        def similarity(a: set[str], b: set[str]) -> float:
            if not a or not b:
                return 0.0
            return len(a & b) / len(a | b)

        tokenized = [tokenize(s) for s in sentences]
        parent_chunks = []
        current_chunk = [sentences[0]]
        current_tokens = set(tokenized[0])
        current_length = len(sentences[0])
        for i in range(1, len(sentences)):
            sim = similarity(current_tokens, tokenized[i])
            if sim < similarity_threshold and current_length > min_chunk_size:
                parent_chunks.append("".join(current_chunk))
                current_chunk = [sentences[i]]
                current_tokens = set(tokenized[i])
                current_length = len(sentences[i])
            else:
                current_chunk.append(sentences[i])
                current_tokens.update(tokenized[i])
                current_length += len(sentences[i])
        if current_chunk:
            parent_chunks.append("".join(current_chunk))
        return parent_chunks
    
    @staticmethod
    def chunk_code(code: str, language_suffix: str) -> list[str]:
        """Heuristic code segment splitter."""
        if not code or not code.strip():
            return []

        lines = code.splitlines()
        chunks = []
        current_chunk = []
        current_length = 0
        
        brace_languages = {
            ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".cc", ".cxx", 
            ".hpp", ".hh", ".hxx", ".cs", ".go", ".rs", ".php", ".swift", ".kt", 
            ".kts", ".scala", ".dart", ".groovy", ".m", ".mm", ".zig", ".vue", 
            ".css", ".scss", ".sass", ".less", ".json", 
        }
        
        brace_count = 0
        for line in lines:
            current_chunk.append(line)
            current_length += len(line)
            
            if language_suffix in brace_languages:
                brace_count += line.count("{") - line.count("}")
                if brace_count == 0 and current_length > 800:
                    chunks.append("\n".join(current_chunk))
                    current_chunk = []
                    current_length = 0
            else:
                if (line.startswith("def ") or line.startswith("class ") or line.startswith("import ")) and current_length > 800:
                    next_start = current_chunk.pop()
                    if current_chunk:
                        chunks.append("\n".join(current_chunk))
                    current_chunk = [next_start]
                    current_length = len(next_start)

        if current_chunk:
            chunks.append("\n".join(current_chunk))
            
        return chunks

if __name__ == "__main__":
    pass